from datetime import date

import pytest
from conftest import ROOT, emp, punches

from hr_agent.devices import load_profile, parse_export
from hr_agent.models import PayrollRun
from hr_agent.storage import RunLockedError


def test_zkteco_csv_profile(tmp_path):
    f = tmp_path / "x.csv"
    f.write_text("AC-No.,Name,Time,State\n101,A,01/09/2026 07:55:00,C/In\n101,A,bad,C/Out\n,B,01/09/2026 08:00:00,C/In\n")
    ps, errors = parse_export(f, load_profile("zkteco_csv", ROOT / "config/device_profiles"))
    assert len(ps) == 1 and ps[0].direction == "in" and ps[0].employee_id == "101"
    assert len(errors) == 2


def test_punches_roundtrip_and_dedupe(store):
    ps = punches("E1", ("2026-09-01T08:00", "2026-09-01T17:00"))
    assert store.save_punches(ps, "a.csv", "hr") == 2
    assert store.save_punches(ps, "a.csv", "hr") == 0
    assert len(store.punches(date(2026, 9, 1), date(2026, 9, 30))) == 2


def test_locked_run_cannot_be_rerun(store):
    run = PayrollRun(factory_id="f", period="2026-09", created_at="2026-10-01T00:00:00Z", created_by="hr", payslips=[], flags=[])
    store.save_run(run)
    for s in ("reviewed", "approved", "locked"):
        store.set_status("2026-09", 1, s, "u")
    assert store.load_run("2026-09").status == "locked"
    with pytest.raises(RunLockedError):
        store.save_run(PayrollRun(factory_id="f", period="2026-09", created_at="2026-10-02T00:00:00Z", created_by="hr", payslips=[], flags=[]))


def test_employee_master_reads_back(store):
    from openpyxl import load_workbook

    from hr_agent.storage.excel import MASTER_SHEETS

    wb = load_workbook(store.master)
    e = emp("E9", spouse_working=False).model_dump()
    e["spouse_working"] = "N"
    wb["Employees"].append([e.get(c) for c in MASTER_SHEETS["Employees"]])
    wb.save(store.master)
    (got,) = store.employees()
    assert got.employee_id == "E9" and got.spouse_working is False and str(got.monthly_salary) == "1700"


def test_pay_items_sheet_reads_back(store):
    from datetime import datetime

    from openpyxl import load_workbook

    from hr_agent.storage.excel import MASTER_SHEETS
    wb = load_workbook(store.master)
    rows = [{"employee_id": 101, "item": "incentive", "amount": 150, "period": datetime(2026, 9, 1), "approved_by": "GM"},
            {"employee_id": "101", "item": "FIXED_ALW", "amount": 300}]
    for r in rows:
        wb["PayItems"].append([r.get(c) for c in MASTER_SHEETS["PayItems"]])
    wb.save(store.master)
    a, b = store.pay_items()
    assert (a.employee_id, a.item, a.period, a.approved_by) == ("101", "INCENTIVE", "2026-09", "GM")
    assert b.period is None and b.applies_to("2026-10", date(2026, 10, 1), date(2026, 10, 31))
