import hashlib
import hmac
import json
from decimal import Decimal

import httpx
import pytest
from conftest import emp
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_end_to_end import TOKENS, _fake_llm

from hr_agent import service
from hr_agent.channels.gateway import TEXTS, ChatGateway, StaffMember
from hr_agent.channels.telegram import TelegramClient
from hr_agent.channels.whatsapp import WhatsAppClient
from hr_agent.cloud_agent.agent import Agent, Tenant
from hr_agent.cloud_agent.app import create_cloud_app
from hr_agent.factory_node.app import create_app
from hr_agent.storage.excel import MASTER_SHEETS

SECRET, APP_SECRET = "tg-secret", "meta-secret"


@pytest.fixture
def chat(factory, store, monkeypatch):
    wb = load_workbook(store.master)
    for e in (emp("F1", phone="012-345 6789"), emp("L1", "3000", "local", telegram_id="555", language="ms")):
        wb["Employees"].append([e.model_dump().get(c) for c in MASTER_SHEETS["Employees"]])
    wb.save(store.master)
    monkeypatch.setenv("TG_SECRET", SECRET)
    sent = {"telegram": [], "whatsapp": []}

    def recorder(kind):
        def handler(request: httpx.Request):
            sent[kind].append(json.loads(request.content))
            return httpx.Response(200, json={"ok": True})
        return httpx.MockTransport(handler)

    calls = []
    tenant = Tenant(factory.id, "http://node", "X", telegram_secret_env="TG_SECRET", whatsapp_phone_number_id="PN1")
    agent = Agent({factory.id: tenant}, llm=_fake_llm(calls), model="m")
    staff = [StaffMember(factory.id, "hr1", "hr", telegram_id="900"), StaffMember(factory.id, "acct", "accountant", whatsapp="60111111111")]
    gateway = ChatGateway(agent, staff, node_client=lambda t: TestClient(
        create_app(factory, store, TOKENS), headers={"Authorization": "Bearer t-agent"}))
    app = create_cloud_app(gateway, telegram=lambda t: TelegramClient("T", transport=recorder("telegram")),
                           whatsapp=lambda t: WhatsAppClient("W", t.whatsapp_phone_number_id, "v21.0", transport=recorder("whatsapp")),
                           whatsapp_app_secret=APP_SECRET, whatsapp_verify_token="verify-me")
    return TestClient(app), sent, calls


def tg(client, update_id, user_id, text=None, contact=None, chat_type="private"):
    msg = {"chat": {"id": int(user_id), "type": chat_type}, "from": {"id": int(user_id), "language_code": "en"}}
    if text is not None:
        msg["text"] = text
    if contact is not None:
        msg["contact"] = contact
    return client.post(f"/telegram/{'demo-pahang'}", json={"update_id": update_id, "message": msg},
                       headers={"X-Telegram-Bot-Api-Secret-Token": SECRET})


def test_telegram_webhook_secret_and_unknown_factory(chat):
    client, sent, _ = chat
    assert client.post("/telegram/demo-pahang", json={}, headers={"X-Telegram-Bot-Api-Secret-Token": "x"}).status_code == 403
    assert client.post("/telegram/nope", json={}, headers={"X-Telegram-Bot-Api-Secret-Token": SECRET}).status_code == 404
    assert tg(client, 1, 777, "hi", chat_type="group").status_code == 200 and sent["telegram"] == []  # groups ignored


def test_telegram_employee_links_by_phone_then_asks_about_own_payslip(chat, store):
    client, sent, calls = chat
    tg(client, 1, 777, "my payslip?")
    assert sent["telegram"][-1]["text"] == TEXTS["unregistered_telegram"]
    assert sent["telegram"][-1]["reply_markup"]["keyboard"][0][0]["request_contact"] is True
    tg(client, 2, 777, contact={"user_id": 888, "phone_number": "+60123456789"})  # someone else's contact
    assert sent["telegram"][-1]["text"] == TEXTS["own_contact"]
    tg(client, 3, 777, contact={"user_id": 777, "phone_number": "+60 19 999 9999"})
    assert sent["telegram"][-1]["text"] == TEXTS["not_matched"]
    tg(client, 4, 777, contact={"user_id": 777, "phone_number": "+60123456789"})
    assert sent["telegram"][-1]["text"] == TEXTS["linked"]
    assert [(lk.employee_id, lk.chat_id) for lk in store.chat_links()] == [("F1", "777")]

    tg(client, 5, 777, "show L1's payslip for 2026-09")
    assert sent["telegram"][-1]["text"] == "done"
    assert [t["function"]["name"] for t in calls[0]["tools"]] == ["get_payslip"]
    tool_result = calls[1]["messages"][-1]["content"]
    assert "L1" not in tool_result and "404" in tool_result  # pinned to F1, and drafts are not released

    tg(client, 5, 777, "duplicate delivery")
    assert len(sent["telegram"]) == 5  # retried update answered once


def test_telegram_staff_roles(chat):
    client, sent, calls = chat
    tg(client, 1, 900, "Record 1 day replacement leave for L1")
    assert "record_replacement_leave" in [t["function"]["name"] for t in calls[0]["tools"]]
    calls.clear()
    tg(client, 2, 555, "/start")  # L1 is matched from telegram_id in the master file
    assert sent["telegram"][-1]["text"] == TEXTS["welcome"] and calls == []


def signed(payload):
    raw = json.dumps(payload).encode()
    sig = "sha256=" + hmac.new(APP_SECRET.encode(), raw, hashlib.sha256).hexdigest()
    return raw, {"X-Hub-Signature-256": sig, "Content-Type": "application/json"}


def wa(msg_id, sender, text=None, kind="text", pnid="PN1"):
    m = {"id": msg_id, "from": sender, "type": kind}
    if text is not None:
        m["text"] = {"body": text}
    return {"entry": [{"changes": [{"value": {"metadata": {"phone_number_id": pnid}, "messages": [m]}}]}]}


def test_whatsapp_verify_and_signature(chat):
    client, sent, _ = chat
    ok = client.get("/whatsapp", params={"hub.mode": "subscribe", "hub.verify_token": "verify-me", "hub.challenge": "42"})
    assert ok.status_code == 200 and ok.text == "42"
    assert client.get("/whatsapp", params={"hub.mode": "subscribe", "hub.verify_token": "x"}).status_code == 403
    raw, headers = signed(wa("m1", "60123456789", "hi"))
    assert client.post("/whatsapp", content=raw, headers={**headers, "X-Hub-Signature-256": "sha256=bad"}).status_code == 403
    assert sent["whatsapp"] == []


def test_whatsapp_employee_staff_and_unknown(chat):
    client, sent, calls = chat
    for msg in (wa("m1", "60123456789", "my OT?"), wa("m2", "60199999999", "hi"), wa("m3", "60123456789", kind="image"),
                wa("m4", "60111111111", "run summary 2026-09"), wa("m5", "60123456789", "x", pnid="OTHER")):
        raw, headers = signed(msg)
        assert client.post("/whatsapp", content=raw, headers=headers).status_code == 200
    assert [(m["to"], m["text"]["body"]) for m in sent["whatsapp"]] == [
        ("60123456789", "done"), ("60199999999", TEXTS["unregistered_whatsapp"]),
        ("60123456789", TEXTS["text_only"]), ("60111111111", "done")]
    assert sent["whatsapp"][0]["messaging_product"] == "whatsapp"
    accountant_tools = [t["function"]["name"] for t in calls[2]["tools"]]
    assert "get_run_summary" in accountant_tools and "record_replacement_leave" not in accountant_tools


def test_node_directory_is_agent_only(factory, store):
    wb = load_workbook(store.master)
    wb["Employees"].append([emp("F1", phone="0123456789").model_dump().get(c) for c in MASTER_SHEETS["Employees"]])
    wb.save(store.master)
    node = TestClient(create_app(factory, store, TOKENS))
    assert node.get("/directory/employee", params={"phone": "60123456789"}, headers={"Authorization": "Bearer t-hr"}).status_code == 403
    r = node.get("/directory/employee", params={"phone": "+60 12-345 6789"}, headers={"Authorization": "Bearer t-agent"})
    assert r.json() == {"employee_id": "F1", "language": "en"}
    assert service.find_employee(store, phone="60199999999") is None


def test_staff_attendance_tool(factory, store, tmp_path):
    wb = load_workbook(store.master)
    for e in (emp("F1", device_user_id="101", department="Line 2"), emp("F2", device_user_id="102", department="Line 3")):
        wb["Employees"].append([e.model_dump().get(c) for c in MASTER_SHEETS["Employees"]])
    wb.save(store.master)
    csv = tmp_path / "att.csv"
    rows = ["AC-No.,Name,Time,State", "101,X,01/09/2026 08:20:00,C/In", "101,X,01/09/2026 19:00:00,C/Out",
            "102,X,01/09/2026 07:55:00,C/In"]
    csv.write_text("\n".join(rows) + "\n")
    service.import_attendance(factory, store, csv, "zkteco_csv", "hr1")
    service.run_payroll(factory, store, "2026-09", "hr1")
    node = TestClient(create_app(factory, store, TOKENS))
    assert node.get("/runs/2026-09/attendance", headers={"Authorization": "Bearer t-f1"}).status_code == 403
    r = node.get("/runs/2026-09/attendance", params={"department": "line 2"}, headers={"Authorization": "Bearer t-gm"}).json()
    assert r["departments"] == ["Line 2", "Line 3"] and [e["employee_id"] for e in r["employees"]] == ["F1"]
    f1 = r["employees"][0]
    assert f1["late_days"] == 1 and f1["late_minutes"] == 20 and Decimal(f1["ot_hours"]) == 2 and f1["absent_days"] > 0

    calls = []
    agent = Agent({factory.id: Tenant(factory.id, "http://node", "X")},
                  llm=_fake_llm(calls, "get_attendance", {"period": "2026-09", "department": "Line 3"}), model="m")
    agent.ask(factory.id, "Who missed punches in Line 3?", http=TestClient(
        create_app(factory, store, TOKENS), headers={"Authorization": "Bearer t-agent"}))
    result = json.loads(calls[1]["messages"][-1]["content"])
    assert [e["employee_id"] for e in result["employees"]] == ["F2"] and result["employees"][0]["missing_punch_days"] == 1
    calls.clear()
    agent.ask(factory.id, "my attendance", employee_id="F1", http=node)
    assert "get_attendance" not in [t["function"]["name"] for t in calls[0]["tools"]]
