"""A device profile describes how one brand/model exports punches.

Adding a new device model means adding a YAML profile, not changing code.
"""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
from typing import Literal

import yaml
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from hr_agent.models import Punch

ColumnRef = str | int


class ColumnMap(BaseModel):
    employee_id: ColumnRef
    timestamp: ColumnRef | None = None
    date: ColumnRef | None = None
    time: ColumnRef | None = None
    direction: ColumnRef | None = None
    device_id: ColumnRef | None = None


class DeviceProfile(BaseModel):
    id: str
    name: str
    file_type: Literal["csv", "xlsx"] = "csv"
    delimiter: str = ","
    encoding: str = "utf-8-sig"
    has_header: bool = True
    skip_rows: int = 0
    sheet: str | None = None
    columns: ColumnMap
    datetime_formats: list[str] = Field(default_factory=lambda: ["%Y-%m-%d %H:%M:%S"])
    date_formats: list[str] = Field(default_factory=lambda: ["%Y-%m-%d"])
    time_formats: list[str] = Field(default_factory=lambda: ["%H:%M:%S", "%H:%M"])
    direction_map: dict[str, Literal["in", "out", "unknown"]] = Field(default_factory=dict)


class ParseError(BaseModel):
    row: int
    message: str


def load_profile(profile_id: str, profiles_dir: str | Path) -> DeviceProfile:
    path = Path(profiles_dir) / f"{profile_id}.yaml"
    if not path.exists():
        raise FileNotFoundError(f"Device profile {profile_id!r} not found in {profiles_dir}")
    return DeviceProfile.model_validate(yaml.safe_load(path.read_text(encoding="utf-8")))


def _read_rows(path: Path, profile: DeviceProfile) -> list[list[str]]:
    if profile.file_type == "xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[profile.sheet] if profile.sheet else wb.worksheets[0]
        rows = [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        delimiter = "\t" if profile.delimiter in ("\\t", "tab") else profile.delimiter
        with path.open(encoding=profile.encoding, newline="") as fh:
            rows = list(csv.reader(fh, delimiter=delimiter))
    return rows[profile.skip_rows :]


def _parse_dt(value, formats: list[str]) -> datetime:
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    raise ValueError(f"unrecognised date/time {text!r}")


def parse_export(
    path: str | Path, profile: DeviceProfile
) -> tuple[list[Punch], list[ParseError]]:
    """Return punches keyed by *device user id* plus per-row errors."""
    rows = _read_rows(Path(path), profile)
    if not rows:
        return [], []
    header: list[str] = []
    if profile.has_header:
        header = [str(h).strip() for h in rows[0]]
        rows = rows[1:]
    first_data_row = profile.skip_rows + (2 if profile.has_header else 1)

    def idx(ref: ColumnRef | None) -> int | None:
        if ref is None:
            return None
        if isinstance(ref, int):
            return ref
        if ref not in header:
            raise ValueError(f"Column {ref!r} not found in export header {header}")
        return header.index(ref)

    cols = profile.columns
    i_emp, i_ts, i_date, i_time = idx(cols.employee_id), idx(cols.timestamp), idx(cols.date), idx(cols.time)
    i_dir, i_dev = idx(cols.direction), idx(cols.device_id)
    if i_ts is None and (i_date is None or i_time is None):
        raise ValueError("Profile must map either 'timestamp' or both 'date' and 'time'")

    punches: list[Punch] = []
    errors: list[ParseError] = []
    for n, row in enumerate(rows, start=first_data_row):
        if not any(str(c).strip() for c in row):
            continue
        try:
            emp = str(row[i_emp]).strip()
            if emp.endswith(".0"):
                emp = emp[:-2]
            if not emp:
                raise ValueError("missing employee id")
            if i_ts is not None:
                ts = _parse_dt(row[i_ts], profile.datetime_formats)
            else:
                d = _parse_dt(row[i_date], profile.date_formats).date()
                tval = row[i_time]
                t = tval.time() if isinstance(tval, datetime) else _parse_dt(tval, profile.time_formats).time()
                ts = datetime.combine(d, t)
            direction = "unknown"
            if i_dir is not None:
                direction = profile.direction_map.get(str(row[i_dir]).strip(), "unknown")
            device = str(row[i_dev]).strip() if i_dev is not None else ""
            punches.append(Punch(employee_id=emp, timestamp=ts, direction=direction, device_id=device))
        except (ValueError, IndexError) as exc:
            errors.append(ParseError(row=n, message=str(exc)))
    return punches, errors
