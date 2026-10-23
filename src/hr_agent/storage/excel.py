"""Excel-only storage (V1). Each factory's data lives in workbooks on its own server.

Layout under the storage root:
  master.xlsx                      Employees, PayItems, Leave, LeaveBalances, LeaveCredits, Overrides, ChatLinks
  attendance/punches_YYYY-MM.xlsx  imported punches, one file per month
  payroll/YYYY-MM/run_vN.xlsx      one immutable file per payroll run version
  audit_log.xlsx                   append-only audit trail
Writes are serialised with a file lock; locked runs are never rewritten.
"""

from __future__ import annotations

import json
from datetime import date, datetime, time, timedelta
from pathlib import Path

from filelock import FileLock
from openpyxl import Workbook, load_workbook

from hr_agent.models import (
    AttendanceOverride,
    ChatLink,
    Employee,
    LeaveBalance,
    LeaveCredit,
    LeaveRecord,
    PayItemEntry,
    PayrollRun,
    Punch,
    RunStatus,
    StatusChange,
)
from hr_agent.storage.base import RunLockedError, Store

MASTER_SHEETS: dict[str, list[str]] = {
    "Employees": list(Employee.model_fields),
    "PayItems": list(PayItemEntry.model_fields),
    "Leave": list(LeaveRecord.model_fields),
    "LeaveBalances": list(LeaveBalance.model_fields),
    "LeaveCredits": list(LeaveCredit.model_fields),
    "Overrides": list(AttendanceOverride.model_fields),
    "ChatLinks": list(ChatLink.model_fields),
}
_BOOL = {"y": True, "yes": True, "true": True, "1": True, "n": False, "no": False, "false": False, "0": False}
_CHUNK = 30000


def _clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        v = value.strip()
        if not v:
            return None
        return _BOOL.get(v.lower(), v) if v.lower() in _BOOL and len(v) <= 5 else v
    if isinstance(value, datetime):
        return value.date() if value.time() == time(0) else value
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return value


def _read_sheet(path: Path, sheet: str) -> list[dict]:
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return []
        rows = list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()
    while rows and all(v is None for v in rows[0]):
        rows.pop(0)
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for row in rows[1:]:
        rec = {h: _clean(v) for h, v in zip(header, row, strict=False) if h}
        rec = {k: v for k, v in rec.items() if v is not None}
        if rec:
            out.append(rec)
    return out


def _as_text_ids(rec: dict, *keys: str) -> dict:
    for k in keys:
        if k in rec:
            rec[k] = str(rec[k])
    return rec


def _cell(v):
    if isinstance(v, (list, dict)):
        return json.dumps(v, default=str)
    if hasattr(v, "is_finite"):  # Decimal
        return float(v)
    return v


def _write_rows(ws, header: list[str], rows: list[dict]) -> None:
    ws.append(header)
    for r in rows:
        ws.append([_cell(r.get(h)) for h in header])


class ExcelStore(Store):
    def __init__(self, root: str | Path):
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)
        self.master = self.root / "master.xlsx"
        self._lock = FileLock(str(self.root / ".write.lock"))

    # Template --------------------------------------------------------------
    def init_master(self) -> Path:
        if self.master.exists():
            return self.master
        wb = Workbook()
        info = wb.active
        info.title = "Instructions"
        info.append(["Fill one row per record. Dates as YYYY-MM-DD; booleans as Y/N."])
        info.append(["PayItems: item is a pay-item code from the factory config; period YYYY-MM for one-off amounts,"
                     " or blank with start/end dates for recurring ones."])
        info.append(["Leave types: " + ", ".join(t for t in LeaveRecord.model_fields["leave_type"].annotation.__members__)])
        for name, header in MASTER_SHEETS.items():
            wb.create_sheet(name).append(header)
        wb.save(self.master)
        return self.master

    # Master data -------------------------------------------------------------
    def employees(self) -> list[Employee]:
        rows = _read_sheet(self.master, "Employees")
        return [Employee.model_validate(_as_text_ids(r, "employee_id", "device_user_id", "bank_account", "phone", "telegram_id", "id_no")) for r in rows]

    def pay_items(self) -> list[PayItemEntry]:
        out = []
        for r in _read_sheet(self.master, "PayItems"):
            if isinstance(r.get("period"), (date, datetime)):  # Excel turns "2026-09" into a date
                r["period"] = f"{r['period']:%Y-%m}"
            r = _as_text_ids(r, "employee_id", "period")
            r["item"] = str(r.get("item", "")).strip().upper()
            out.append(PayItemEntry.model_validate(r))
        return out

    def leave(self, start: date, end: date) -> list[LeaveRecord]:
        recs = [LeaveRecord.model_validate(_as_text_ids(r, "employee_id")) for r in _read_sheet(self.master, "Leave")]
        return [r for r in recs if r.start <= end and r.end >= start]

    def leave_balances(self, year: int) -> list[LeaveBalance]:
        recs = [LeaveBalance.model_validate(_as_text_ids(r, "employee_id")) for r in _read_sheet(self.master, "LeaveBalances")]
        return [r for r in recs if r.year == year]

    def leave_credits(self, year: int) -> list[LeaveCredit]:
        recs = [LeaveCredit.model_validate(_as_text_ids(r, "employee_id", "manager_confirmed_by"))
                for r in _read_sheet(self.master, "LeaveCredits")]
        return [r for r in recs if r.year == year]

    def _append_master(self, sheet: str, record) -> None:
        header = MASTER_SHEETS[sheet]
        with self._lock:
            wb = load_workbook(self.master)
            ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
            if ws.max_row == 1 and all(c.value is None for c in ws[1]):
                ws.append(header)
            # Follow the sheet's own column order (HR may reorder columns); add any missing columns at the end.
            existing = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
            for h in header:
                if h not in existing:
                    existing.append(h)
                    ws.cell(1, len(existing)).value = h
            row = record.model_dump()
            ws.append([_cell(row[h].value if hasattr(row[h], "value") else row[h]) if h in row else None for h in existing])
            wb.save(self.master)

    def chat_links(self) -> list[ChatLink]:
        return [ChatLink.model_validate(_as_text_ids(r, "employee_id", "chat_id", "phone")) for r in _read_sheet(self.master, "ChatLinks")]

    def add_chat_link(self, link: ChatLink) -> ChatLink:
        self._append_master("ChatLinks", link)
        self.audit("chat_link", link.linked_by, f"{link.employee_id} linked {link.channel} account")
        return link

    def add_leave_credit(self, credit: LeaveCredit) -> LeaveCredit:
        with self._lock:  # duplicate check and append under one lock (FileLock is re-entrant)
            if credit.worked_on and any(c.employee_id == credit.employee_id and c.leave_type == credit.leave_type
                                        and c.worked_on == credit.worked_on for c in self.leave_credits(credit.year)):
                raise ValueError(f"{credit.leave_type.value} leave for {credit.worked_on} is already recorded "
                                 f"for {credit.employee_id}")
            self._append_master("LeaveCredits", credit)
        who = credit.recorded_by + (f" for {credit.on_behalf_of}" if credit.on_behalf_of else "")
        self.audit("leave_credit", who, f"{credit.employee_id} +{credit.days} {credit.leave_type.value} {credit.year}"
                   f" (worked {credit.worked_on}; manager {credit.manager_confirmed_by}; {credit.reason})")
        return credit

    def overrides(self, start: date, end: date) -> list[AttendanceOverride]:
        recs = [AttendanceOverride.model_validate(_as_text_ids(r, "employee_id")) for r in _read_sheet(self.master, "Overrides")]
        return [r for r in recs if start <= r.date <= end]

    # Punches -----------------------------------------------------------------
    def _punch_file(self, year: int, month: int) -> Path:
        return self.root / "attendance" / f"punches_{year:04d}-{month:02d}.xlsx"

    def save_punches(self, punches: list[Punch], source: str, user: str) -> int:
        by_month: dict[tuple[int, int], list[Punch]] = {}
        for p in punches:
            by_month.setdefault((p.timestamp.year, p.timestamp.month), []).append(p)
        added = 0
        with self._lock:
            for (y, m), items in by_month.items():
                path = self._punch_file(y, m)
                path.parent.mkdir(parents=True, exist_ok=True)
                existing = {(r["employee_id"], r["timestamp"]) for r in self._punch_rows(path)}
                if path.exists():
                    wb = load_workbook(path)
                    ws = wb["Punches"]
                else:
                    wb = Workbook()
                    ws = wb.active
                    ws.title = "Punches"
                    ws.append(["employee_id", "timestamp", "direction", "device_id", "source", "imported_by", "imported_at"])
                now = datetime.now().replace(microsecond=0)
                for p in sorted(items, key=lambda x: (x.employee_id, x.timestamp)):
                    key = (p.employee_id, p.timestamp.replace(microsecond=0))
                    if key in existing:
                        continue
                    existing.add(key)
                    ws.append([p.employee_id, key[1], p.direction, p.device_id, source, user, now])
                    added += 1
                wb.save(path)
        self.audit("import_punches", user, f"{added} new punches from {source}")
        return added

    @staticmethod
    def _punch_rows(path: Path) -> list[dict]:
        rows = _read_sheet(path, "Punches")
        for r in rows:
            r["employee_id"] = str(r["employee_id"])
            ts = r["timestamp"]
            r["timestamp"] = ts if isinstance(ts, datetime) else datetime.combine(ts, time(0)) if isinstance(ts, date) else datetime.fromisoformat(str(ts))
        return rows

    def punches(self, start: date, end: date) -> list[Punch]:
        out: list[Punch] = []
        y, m = start.year, start.month
        while (y, m) <= (end.year, end.month):
            for r in self._punch_rows(self._punch_file(y, m)):
                if start <= r["timestamp"].date() <= end:
                    out.append(Punch(employee_id=r["employee_id"], timestamp=r["timestamp"], direction=r.get("direction", "unknown"), device_id=str(r.get("device_id") or "")))
            y, m = (y + 1, 1) if m == 12 else (y, m + 1)
        return out

    # Payroll runs -------------------------------------------------------------
    def _run_dir(self, period: str) -> Path:
        return self.root / "payroll" / period

    def _versions(self, period: str) -> list[int]:
        d = self._run_dir(period)
        if not d.exists():
            return []
        return sorted(int(p.stem.split("_v")[1]) for p in d.glob("run_v*.xlsx"))

    def run_path(self, period: str, version: int) -> Path:
        return self._run_dir(period) / f"run_v{version}.xlsx"

    def _write_run(self, run: PayrollRun) -> None:
        path = self.run_path(run.period, run.version)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        meta = wb.active
        meta.title = "Summary"
        meta.append(["Factory", run.factory_id])
        meta.append(["Period", run.period])
        meta.append(["Version", run.version])
        meta.append(["Status", run.status])
        meta.append(["Created", run.created_at.replace(tzinfo=None), run.created_by])
        meta.append([])
        header = ["employee_id", "name", "gross", "total_deductions", "net", "employer_cost", "ot_hours", "epf_wage", "socso_wage"]
        meta.append(header)
        for p in run.payslips:
            meta.append([p.employee_id, p.employee_name, float(p.gross), float(p.total_deductions), float(p.net), float(p.employer_cost), float(p.ot_hours_total), float(p.epf_wage), float(p.socso_wage)])
        meta.append(["TOTAL", "", *[float(sum(getattr(p, f) for p in run.payslips)) for f in ("gross", "total_deductions", "net", "employer_cost", "ot_hours_total")]])

        lines = wb.create_sheet("Lines")
        lines.append(["employee_id", "kind", "code", "label", "quantity", "rate", "amount", "trace"])
        for p in run.payslips:
            for ln in p.lines:
                lines.append([p.employee_id, ln.kind, ln.code, ln.label, _cell(ln.quantity), _cell(ln.rate), float(ln.amount), ln.trace])
        days = wb.create_sheet("Attendance")
        dh = ["employee_id", "date", "day_type", "status", "check_in", "check_out", "normal_hours", "ot_hours", "worked_hours", "late_minutes", "leave_type", "overridden"]
        days.append(dh)
        for p in run.payslips:
            for dr in p.days:
                row = dr.model_dump()
                days.append([p.employee_id, *[_cell(row[h].value if hasattr(row[h], "value") else row[h]) for h in dh[1:]]])
        flags = wb.create_sheet("Flags")
        _write_rows(flags, ["severity", "code", "employee_id", "date", "message"], [f.model_dump() for f in run.flags])
        hist = wb.create_sheet("History")
        _write_rows(hist, ["status", "by", "at", "note"], [{**h.model_dump(), "at": h.at.replace(tzinfo=None)} for h in run.history])
        data = wb.create_sheet("_data")
        blob = run.model_dump_json()
        for i in range(0, len(blob), _CHUNK):
            data.append([blob[i : i + _CHUNK]])
        data.sheet_state = "hidden"
        tmp = path.with_suffix(".tmp.xlsx")
        wb.save(tmp)
        tmp.replace(path)

    def save_run(self, run: PayrollRun) -> PayrollRun:
        with self._lock:
            versions = self._versions(run.period)
            if versions:
                latest = self._load(run.period, versions[-1])
                if latest and latest.status == "locked":
                    raise RunLockedError(f"Payroll {run.period} is locked (v{latest.version})")
            run.version = (versions[-1] + 1) if versions else 1
            self._write_run(run)
        self.audit("payroll_run", run.created_by, f"{run.period} v{run.version}")
        return run

    def _load(self, period: str, version: int) -> PayrollRun | None:
        path = self.run_path(period, version)
        if not path.exists():
            return None
        wb = load_workbook(path, read_only=True)
        try:
            blob = "".join(r[0] or "" for r in wb["_data"].iter_rows(values_only=True))
        finally:
            wb.close()
        return PayrollRun.model_validate_json(blob)

    def load_run(self, period: str, version: int | None = None) -> PayrollRun | None:
        versions = self._versions(period)
        if not versions:
            return None
        return self._load(period, version or versions[-1])

    def runs_for_year(self, year: int, locked_only: bool = True) -> list[PayrollRun]:
        base = self.root / "payroll"
        runs = []
        if not base.exists():
            return runs
        for d in sorted(base.glob(f"{year:04d}-*")):
            run = self.load_run(d.name)
            if run and (run.status == "locked" or not locked_only):
                runs.append(run)
        return runs

    def set_status(self, period: str, version: int, status: RunStatus, user: str, note: str = "") -> PayrollRun:
        with self._lock:
            run = self._load(period, version)
            if run is None:
                raise FileNotFoundError(f"No payroll run {period} v{version}")
            if run.status == "locked":
                raise RunLockedError(f"Payroll {period} v{version} is locked")
            run.status = status
            run.history.append(StatusChange(status=status, by=user, at=datetime.now().replace(microsecond=0), note=note))
            self._write_run(run)
        self.audit(f"payroll_{status}", user, f"{period} v{version} {note}".strip())
        return run

    # Audit ---------------------------------------------------------------------
    def audit(self, action: str, user: str, details: str = "") -> None:
        path = self.root / "audit_log.xlsx"
        lock = FileLock(str(self.root / ".audit.lock"))
        with lock:
            if path.exists():
                wb = load_workbook(path)
                ws = wb.active
            else:
                wb = Workbook()
                ws = wb.active
                ws.title = "Audit"
                ws.append(["timestamp", "user", "action", "details"])
            ws.append([datetime.now().replace(microsecond=0), user, action, details])
            wb.save(path)


def punch_window(start: date, end: date) -> tuple[date, date]:
    """Punches needed to cover a window, including after-midnight punches."""
    return start, end + timedelta(days=1)
